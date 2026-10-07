import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_sla_excel(data: dict = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "SLA por Polo"

    title_font = Font(name="Calibri", size=16, bold=True, color="002D62")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="002D62", end_color="002D62", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    green_font = Font(name="Calibri", size=10, bold=True, color="16A34A")
    yellow_font = Font(name="Calibri", size=10, bold=True, color="CA8A04")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws["A1"] = "OTIS SMARTFLOW IA — GESTÃO DE DISPONIBILIDADE E SLA"
    ws["A1"].font = title_font
    ws["A2"] = "Relatório Oficial de Cumprimento de SLA por Polo Regional & Desempenho Operacional"
    ws["A2"].font = subtitle_font

    headers = [
        "Polo Regional", "Supervisor Responsável", "Técnicos Ativos", 
        "Equipamentos", "Chamados Mês", "SLA Global (%)", 
        "Tempo Médio Desloc. (TA)", "Tempo Médio Solução (TB)", "Status SLA"
    ]

    row_num = 4
    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[row_num].height = 26

    rows_data = [
        ("São Paulo - Capital & ABC", "Carlos Silveira", 42, 1850, 542, 98.6, "28 min", "45 min", "CONFORME"),
        ("Rio de Janeiro - Metropolitana", "Marcos Vinicius", 26, 920, 315, 97.8, "34 min", "52 min", "CONFORME"),
        ("Belo Horizonte & Sul de MG", "Renata Figueiredo", 18, 640, 198, 98.9, "25 min", "41 min", "CONFORME"),
        ("Curitiba & Região Sul", "Luciano Brandt", 22, 710, 224, 99.1, "22 min", "38 min", "CONFORME"),
        ("Brasília & Centro-Oeste", "Diego Albuquerque", 14, 430, 121, 96.5, "39 min", "58 min", "ATENÇÃO"),
        ("Nordeste (SSA / REC / FOR)", "Fabio Menezes", 19, 580, 182, 97.2, "32 min", "49 min", "CONFORME"),
        ("Campinas & Interior Paulista", "Guilherme Prado", 16, 510, 147, 98.4, "27 min", "44 min", "CONFORME")
    ]

    for item in rows_data:
        row_num += 1
        ws.row_dimensions[row_num].height = 20
        is_zebra = (row_num % 2 == 0)

        for col_num, val in enumerate(item, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if is_zebra:
                cell.fill = zebra_fill

            if col_num in (3, 4, 5):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_num in (6, 7, 8, 9):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            if col_num == 6:
                cell.value = f"{val}%"
                cell.font = bold_font
            elif col_num == 9:
                cell.font = green_font if val == "CONFORME" else yellow_font

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
