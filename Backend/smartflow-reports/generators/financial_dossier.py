import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_financial_excel(data: dict = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DRE por Contrato"

    title_font = Font(name="Calibri", size=16, bold=True, color="002D62")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="002D62", end_color="002D62", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    green_font = Font(name="Calibri", size=10, bold=True, color="16A34A")
    red_font = Font(name="Calibri", size=10, bold=True, color="DC2626")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws["A1"] = "OTIS SMARTFLOW IA — DOSSIÊ FINANCEIRO & DRE DE CONTRATOS"
    ws["A1"].font = title_font
    ws["A2"] = "Análise de Rentabilidade, Consumo de Peças, Mão de Obra e Margem Operacional por Cliente"
    ws["A2"].font = subtitle_font

    headers = [
        "ID Contrato", "Cliente / Edifício", "Região", "Equipamentos", 
        "Receita Mensal (R$)", "Custo Peças (R$)", "Custo MDO/Desloc. (R$)", 
        "Margem Bruta (R$)", "Margem (%)", "Status Financeiro"
    ]

    row_num = 4
    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[row_num].height = 26

    contracts_data = [
        ("CTR-2024-001", "Condomínio Faria Lima Tower", "SP", 12, 48000.0, 4200.0, 7800.0, 36000.0, 75.0, "LUCRATIVO"),
        ("CTR-2024-004", "Shopping Eldorado Corporate", "SP", 24, 96000.0, 11500.0, 14200.0, 70300.0, 73.2, "LUCRATIVO"),
        ("CTR-2023-089", "Centro Empresarial Botafogo", "RJ", 8, 32000.0, 2900.0, 5100.0, 24000.0, 75.0, "LUCRATIVO"),
        ("CTR-2023-045", "Hospital Mater Dei Contorno", "MG", 10, 45000.0, 9800.0, 8900.0, 26300.0, 58.4, "MARGEM MEDIA"),
        ("CTR-2021-034", "Edifício Comercial Paulista Plaza", "SP", 6, 21000.0, 8200.0, 5800.0, 7000.0, 33.3, "ALERTA MARGEM")
    ]

    for item in contracts_data:
        row_num += 1
        ws.row_dimensions[row_num].height = 20
        is_zebra = (row_num % 2 == 0)

        for col_num, val in enumerate(item, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if is_zebra:
                cell.fill = zebra_fill

            if col_num in (5, 6, 7, 8):
                cell.number_format = 'R$ #,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_num in (4,):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_num in (1, 3, 9, 10):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            if col_num == 9:
                cell.value = f"{val}%"
                cell.font = bold_font
            elif col_num == 10:
                cell.font = green_font if val == "LUCRATIVO" else red_font

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
